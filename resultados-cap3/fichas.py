"""Genera los instrumentos de recoleccion de datos de campo (Word para anexos + Excel para vaciado)."""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

OUT = Path("D:/AppWebPlagasPaltaHass")
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(0)
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)


def p(texto="", negrita=False, cursiva=False, centro=False, doble=True, sangria=False, size=12, after=0):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE if doble else WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(after)
    if sangria:
        pf.first_line_indent = Cm(1.25)
    if centro:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # **negrita** en linea
    partes = texto.split("**")
    for i, t in enumerate(partes):
        if not t:
            continue
        r = par.add_run(t)
        r.bold = negrita or i % 2 == 1
        r.italic = cursiva
        r.font.size = Pt(size)
    return par


def vineta(texto):
    par = p("", doble=True)
    par.paragraph_format.left_indent = Cm(0.9)
    par.paragraph_format.first_line_indent = Cm(-0.5)
    partes = ("•\t" + texto).split("**")
    for i, t in enumerate(partes):
        if t:
            r = par.add_run(t); r.bold = i % 2 == 1
    par.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
    return par


def borde(celda, **lados):
    tcPr = celda._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for lado, val in lados.items():
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), val); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "auto")
        b.append(e)


def sombra(celda, color="E7EFE9"):
    tcPr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def tabla(cab, filas, anchos, size=9, alto=None, grilla=True, cab_rep=True):
    t = doc.add_table(rows=1 + len(filas), cols=len(cab))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, c in enumerate(cab):
        cel = t.rows[0].cells[j]
        cel.text = ""
        r = cel.paragraphs[0].add_run(c); r.bold = True; r.font.size = Pt(size)
        cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        sombra(cel)
    for i, f in enumerate(filas):
        for j, v in enumerate(f):
            cel = t.rows[i + 1].cells[j]
            cel.text = ""
            r = cel.paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
            if j == 0 or len(str(v)) < 6:
                cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        if alto:
            row.height = Cm(alto)
        for j, cel in enumerate(row.cells):
            cel.width = Cm(anchos[j])
            cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for par in cel.paragraphs:
                par.paragraph_format.line_spacing = 1.0
            v = "single" if grilla else "nil"
            borde(cel, top="single", bottom="single", left=v, right=v)
    if cab_rep:
        trPr = t.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    return t


def campos(pares, anchos=(5.5, 10.5)):
    """Bloque de datos generales: etiqueta | linea para completar."""
    t = doc.add_table(rows=len(pares), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (a, b) in enumerate(pares):
        c0, c1 = t.rows[i].cells
        c0.text = ""; c1.text = ""
        r = c0.paragraphs[0].add_run(a); r.bold = True; r.font.size = Pt(10)
        r = c1.paragraphs[0].add_run(b); r.font.size = Pt(10)
        c0.width, c1.width = Cm(anchos[0]), Cm(anchos[1])
        for c in (c0, c1):
            borde(c, top="single", bottom="single", left="single", right="single")
        sombra(c0)
        t.rows[i].height = Cm(0.55)


def titulo_instr(n, nombre):
    p(f"Anexo {n}", negrita=True, doble=False, after=2)
    p(nombre, cursiva=True, doble=False, after=6)


def nueva_seccion(horizontal):
    s = doc.add_section()
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.0 if horizontal else 2.54)


def firmas(cargos):
    p("", doble=False, after=4)
    t = doc.add_table(rows=2, cols=len(cargos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(cargos):
        a, b = t.rows[0].cells[j], t.rows[1].cells[j]
        a.text = "\n_______________________________"; b.text = ""
        r = b.paragraphs[0].add_run(c); r.font.size = Pt(10)
        for x in (a, b):
            x.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


GEN = [("Fundo / unidad agrícola:", "______________________________  Localidad: Chao, Virú, La Libertad"),
       ("Sector:", "☐ A     ☐ B          Fecha: ____ / ____ / 2026"),
       ("Etapa:", "☐ Preprueba (sin sistema)     ☐ Posprueba (con sistema)"),
       ("Evaluador:", "______________________________  Cargo: ________________"),
       ("Condición del día:", "☐ Soleado   ☐ Nublado   ☐ Garúa      Hora de inicio de jornada: ____:____"),
       ("Hoja N.°:", "____ de ____")]

# =====================================================================================  PROTOCOLO
p("Protocolo de aplicación de los instrumentos de recolección de datos en campo", negrita=True, centro=True)
p("Este protocolo establece el tamaño de la muestra, el número de órganos por árbol y el procedimiento que se seguirá "
  "para aplicar los instrumentos en la preprueba y en la posprueba, de modo que ambas mediciones sean comparables.",
  sangria=True)
p("Tamaño de la muestra y unidades de observación", negrita=True)
tabla(["Elemento", "Cantidad", "Detalle"], [
    ["Árboles evaluados", "100", "50 del sector A y 50 del sector B, seleccionados por sorteo del padrón (códigos A-01 a A-50 y B-01 a B-50). Los mismos árboles se evalúan en la preprueba y en la posprueba."],
    ["Órganos por árbol", "10", "5 hojas y 5 frutos: uno por cada orientación de la copa (norte, sur, este, oeste) y uno del interior, a una altura de 1.0 a 2.0 m."],
    ["Imágenes en la posprueba", "1,000", "10 imágenes por árbol (una por órgano), capturadas con la aplicación en el modo de análisis por imagen."],
    ["Frutos para el registro de descarte", "60 por árbol", "6,000 frutos por etapa. Se clasifican al momento de la cosecha o de la selección en campo."],
    ["Imágenes para la concordancia con el especialista", "200", "2 por árbol (1 hoja y 1 fruto, elegidos al azar entre los 10 capturados). Se incluyen además todas las imágenes con alerta de la inspección piloto."],
    ["Evaluadores", "1 técnico + 1 especialista agrónomo", "El mismo técnico realiza la preprueba y la posprueba; el especialista valida los diagnósticos sin conocer el resultado del sistema."],
], [4.0, 2.6, 9.4], size=9)
p("")
p("Procedimiento", negrita=True)
for t_ in [
    "**Preprueba (sin sistema):** el técnico inspecciona visualmente los 10 órganos de cada árbol siguiendo el orden norte, "
    "sur, este, oeste e interior. Se cronometra desde que llega al árbol hasta que termina de anotar el diagnóstico (Ficha 1).",
    "**Tiempo hasta el diagnóstico (preprueba):** cuando el personal del fundo observa un síntoma, se registra la fecha y hora "
    "del aviso y la fecha y hora en que el técnico emite el diagnóstico (Ficha 1, columnas de respuesta).",
    "**Posprueba (con sistema):** en los mismos árboles y órganos, el técnico o el agricultor captura las 10 imágenes con la "
    "aplicación. Se cronometra el mismo intervalo y se anota el código de inspección del historial del sistema (Ficha 2).",
    "**Tiempo hasta el diagnóstico (posprueba):** se registra la hora de captura y la hora en que el técnico confirma el "
    "diagnóstico a distancia (Ficha 2).",
    "**Registro de descarte:** en cada etapa se clasifican 60 frutos por árbol en aptos y descartados, y se anota la causa "
    "fitosanitaria del descarte (Ficha 3).",
    "**Concordancia:** el especialista revisa las 200 imágenes seleccionadas y registra su diagnóstico antes de ver el del "
    "sistema (Ficha 4).",
    "**Control de calidad:** se usa el mismo celular, el mismo cronómetro y el mismo orden de recorrido; se repite la captura "
    "de las imágenes desenfocadas, sobreexpuestas o en las que el órgano seleccionado en la aplicación no corresponda al "
    "fotografiado.",
]:
    vineta(t_)
p("Los datos de las fichas se transcriben a la plantilla de vaciado (archivo Excel), cuyas columnas coinciden con las de "
  "cada ficha.", sangria=True)

# =====================================================================================  FICHA 1
nueva_seccion(True)
titulo_instr("__", "Ficha de observación tecnológica N.° 1: Tiempo de inspección manual (preprueba)")
campos(GEN, (5.0, 20.5))
p("", doble=False, after=6)
filas = [[str(i + 1)] + [""] * 11 for i in range(8)]
tabla(["N.°", "Código del árbol", "Hora de inicio (hh:mm:ss)", "Hora de fin (hh:mm:ss)", "Duración (min)",
       "Hojas revisadas (n)", "Frutos revisados (n)", "Diagnóstico visual del técnico",
       "Fecha y hora del aviso del síntoma", "Fecha y hora del diagnóstico", "Tiempo hasta el diagnóstico (h)",
       "Observaciones"], filas, [0.8, 1.8, 2.0, 2.0, 1.6, 1.6, 1.6, 4.0, 2.6, 2.6, 2.0, 3.1], size=8, alto=0.75)
p("Diagnóstico visual: S = sano; AH = antracnosis en hoja; PL = plagas (trips/ácaros); DN = deficiencia nutricional; "
  "AF = antracnosis en fruto; CE = cercospora; RO = roña; PP = pudrición peduncular; SB = sunblotch. Se pueden anotar varios "
  "códigos.", doble=False, size=9, after=0)
firmas(["Evaluador (técnico)", "Responsable de la investigación"])

# =====================================================================================  FICHA 2
nueva_seccion(True)
titulo_instr("__", "Ficha de observación tecnológica N.° 2: Inspección con el sistema de visión computacional (posprueba)")
campos(GEN[:2] + GEN[3:], (5.0, 20.5))
p("", doble=False, after=6)
tabla(["N.°", "Código del árbol", "Hora de inicio", "Hora de fin", "Duración (min)", "Imágenes de hojas (n)",
       "Imágenes de frutos (n)", "Código(s) de inspección en el sistema", "Diagnóstico del sistema (clase y confianza)",
       "Veredicto (sano / alerta)", "Hora de envío al técnico", "Hora de confirmación", "Tiempo hasta el diagnóstico (h)",
       "Observaciones"], [[str(i + 1)] + [""] * 13 for i in range(8)],
      [0.8, 1.6, 1.5, 1.5, 1.4, 1.4, 1.4, 2.4, 3.4, 1.6, 1.6, 1.6, 1.7, 2.4], size=8, alto=0.75)
p("Diagnóstico del sistema: anotar la clase y la confianza de cada detección mostrada por la aplicación (por ejemplo, "
  "AF 82.7 %). Registrar como observación si el órgano seleccionado no corresponde a la imagen y repetir la captura.",
  doble=False, size=9)
firmas(["Evaluador (técnico)", "Responsable de la investigación"])

# =====================================================================================  FICHA 3
nueva_seccion(True)
titulo_instr("__", "Ficha de observación tecnológica N.° 3: Registro de descarte de frutos por árbol")
campos(GEN[:4] + [("Momento de la evaluación:", "☐ Cosecha   ☐ Selección en campo   ☐ Acopio      Kg cosechados del árbol: ________")],
       (5.0, 20.5))
p("", doble=False, after=6)
tabla(["N.°", "Código del árbol", "Frutos evaluados", "Frutos aptos", "Antracnosis", "Roña", "Cercospora",
       "Pudrición peduncular", "Sunblotch", "Otras causas (plagas, golpes, calibre)", "Total descartados",
       "Descarte (%)", "Kg cosechados", "Observaciones"],
      [[str(i + 1), "", "60"] + [""] * 11 for i in range(8)],
      [0.8, 1.7, 1.5, 1.5, 1.6, 1.4, 1.6, 1.7, 1.6, 2.6, 1.6, 1.5, 1.5, 3.1], size=8, alto=0.75)
p("Total descartados = suma de las causas; Descarte (%) = total descartados / frutos evaluados × 100. El descarte "
  "fitosanitario excluye la columna «Otras causas». Si un fruto presenta más de una afección, se registra en la causa "
  "predominante.", doble=False, size=9)
firmas(["Evaluador (técnico)", "Responsable de la investigación"])

# =====================================================================================  FICHA 4
nueva_seccion(True)
titulo_instr("__", "Ficha de observación tecnológica N.° 4: Concordancia del diagnóstico del sistema con el especialista")
campos([("Especialista:", "______________________________  Profesión / colegiatura: ________________"),
        ("Fecha de revisión:", "____ / ____ / 2026          Hoja N.°: ____ de ____"),
        ("Modalidad:", "☐ Revisión en campo   ☐ Revisión de imágenes del historial del sistema")], (5.0, 20.5))
p("", doble=False, after=6)
tabla(["N.°", "Código de inspección", "Código del árbol", "Órgano (H / F)",
       "Diagnóstico del especialista (antes de ver el sistema)", "Severidad (inicial / avanzada)",
       "Diagnóstico del sistema", "Confianza (%)", "¿Coincide? (Sí / No)", "Acción recomendada por el especialista",
       "Observaciones"], [[str(i + 1)] + [""] * 10 for i in range(9)],
      [0.8, 2.2, 1.7, 1.4, 4.0, 2.0, 3.0, 1.5, 1.6, 3.8, 3.7], size=8, alto=0.75)
p("Coincide = el diagnóstico del sistema y el del especialista corresponden a la misma condición (sano o la misma "
  "enfermedad o plaga). Con esta ficha se calculan la exactitud, la sensibilidad, la especificidad y el coeficiente kappa "
  "de Cohen entre el sistema y el especialista.", doble=False, size=9)
firmas(["Especialista agrónomo", "Responsable de la investigación"])

# =====================================================================================  CUESTIONARIO
nueva_seccion(False)
titulo_instr("__", "Cuestionario de evaluación del nivel de control y reducción de descarte")
campos([("Nombre (opcional):", ""), ("Cargo / función en el fundo:", ""), ("Años de experiencia en palta Hass:", ""),
        ("Fecha:", "____ / ____ / 2026")], (6.0, 10.0))
p("", doble=False, after=6)
p("Instrucciones: después de usar el sistema durante la posprueba, marque con una X el nivel de acuerdo con cada "
  "afirmación. 1 = Totalmente en desacuerdo; 2 = En desacuerdo; 3 = Ni de acuerdo ni en desacuerdo; 4 = De acuerdo; "
  "5 = Totalmente de acuerdo.", doble=False, size=10, after=6)
items = [
    "El sistema permite identificar enfermedades en hojas y frutos antes de que sean evidentes en la inspección habitual.",
    "El diagnóstico del sistema coincide con lo que observo en campo.",
    "El sistema reduce el tiempo que espero para contar con un diagnóstico.",
    "Con el sistema puedo inspeccionar más árboles en una jornada.",
    "Las alertas del sistema me permiten aplicar medidas de control oportunas.",
    "La detección temprana de antracnosis con el sistema reduce los frutos descartados.",
    "El uso del sistema mejora la proporción de frutos aptos para exportación.",
    "El historial y el reporte PDF facilitan el seguimiento sanitario del huerto.",
    "El sistema es fácil de usar en condiciones de campo.",
    "Recomendaría el uso del sistema para la inspección fitosanitaria del fundo.",
]
tabla(["N.°", "Afirmación", "1", "2", "3", "4", "5"], [[str(i + 1), it, "", "", "", "", ""] for i, it in enumerate(items)],
      [1.0, 10.0, 1.0, 1.0, 1.0, 1.0, 1.0], size=10, alto=0.9)
p("", doble=False, after=6)
p("Comentarios u observaciones: ____________________________________________________________________________"
  "______________________________________________________________________", doble=True, size=10)

f = OUT / "Instrumentos de recoleccion de datos - fichas de campo.docx"
doc.save(f)
print("word:", f)

# =====================================================================================  EXCEL DE VACIADO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
cod = [f"A-{i:02d}" for i in range(1, 51)] + [f"B-{i:02d}" for i in range(1, 51)]
FILL = PatternFill("solid", fgColor="E7EFE9")


def hoja(nombre, cab, filas_fn, primera=False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = nombre
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, cd in enumerate(cod):
        ws.append(filas_fn(i + 2, cd))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
    ws.freeze_panes = "C2"
    return ws


hoja("F1_preprueba", ["Código", "Sector", "Fecha", "Hora inicio", "Hora fin", "Duración (min)", "Hojas revisadas",
                      "Frutos revisados", "Diagnóstico visual", "Aviso del síntoma (fecha y hora)",
                      "Diagnóstico (fecha y hora)", "Tiempo hasta diagnóstico (h)", "Observaciones"],
     lambda r, c: [c, c[0], None, None, None, f'=IF(AND(D{r}<>"",E{r}<>""),ROUND((E{r}-D{r})*1440,1),"")', 5, 5, None,
                   None, None, f'=IF(AND(J{r}<>"",K{r}<>""),ROUND((K{r}-J{r})*24,1),"")', None], primera=True)
hoja("F2_posprueba", ["Código", "Sector", "Fecha", "Hora inicio", "Hora fin", "Duración (min)", "Imágenes hojas",
                      "Imágenes frutos", "Códigos de inspección", "Diagnóstico del sistema", "Veredicto",
                      "Envío al técnico (fecha y hora)", "Confirmación (fecha y hora)", "Tiempo hasta diagnóstico (h)",
                      "Observaciones"],
     lambda r, c: [c, c[0], None, None, None, f'=IF(AND(D{r}<>"",E{r}<>""),ROUND((E{r}-D{r})*1440,1),"")', 5, 5, None,
                   None, None, None, None, f'=IF(AND(L{r}<>"",M{r}<>""),ROUND((M{r}-L{r})*24,1),"")', None])
for etapa in ("pre", "post"):
    hoja(f"F3_descarte_{etapa}", ["Código", "Sector", "Fecha", "Frutos evaluados", "Aptos", "Antracnosis", "Roña",
                                  "Cercospora", "Pudrición peduncular", "Sunblotch", "Otras causas", "Total descartados",
                                  "Descarte (%)", "Descarte fitosanitario (%)", "Kg cosechados", "Observaciones"],
         lambda r, c: [c, c[0], None, 60, None, None, None, None, None, None, None, f"=SUM(F{r}:K{r})",
                       f'=IF(D{r}>0,ROUND(L{r}/D{r}*100,1),"")', f'=IF(D{r}>0,ROUND(SUM(F{r}:J{r})/D{r}*100,1),"")',
                       None, None])
ws = wb.create_sheet("F4_concordancia")
ws.append(["N.°", "Código de inspección", "Código del árbol", "Órgano (H/F)", "Diagnóstico especialista", "Severidad",
           "Diagnóstico sistema", "Confianza (%)", "Coincide (1/0)", "Acción recomendada", "Observaciones"])
for c in ws[1]:
    c.font = Font(bold=True); c.fill = FILL; c.alignment = Alignment(wrap_text=True)
for i in range(200):
    ws.append([i + 1] + [None] * 10)
for col in ws.columns:
    ws.column_dimensions[col[0].column_letter].width = 18
dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True); ws.add_data_validation(dv); dv.add("I2:I201")
ws = wb.create_sheet("cuestionario")
ws.append(["Encuestado", "Cargo"] + [f"Ítem {i}" for i in range(1, 11)] + ["Promedio"])
for c in ws[1]:
    c.font = Font(bold=True); c.fill = FILL
for i in range(2, 12):
    ws.append([f"E{i - 1:02d}", None] + [None] * 10 + [f'=IF(COUNT(C{i}:L{i})>0,AVERAGE(C{i}:L{i}),"")'])
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
ws.add_data_validation(dv); dv.add("C2:L11")
fx = OUT / "resultados-cap3/simulacion/plantilla_vaciado_fichas_campo.xlsx"
wb.save(fx)
print("excel:", fx)
